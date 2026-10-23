"""Generate a faculty-safe SysSense AI Review 3 presentation.

The deck distinguishes direct live telemetry from dataset-trained forecasts and
uses the current reproducible benchmark results. It intentionally avoids
unsupported "accuracy" claims for regression and anomaly detection.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "SysSense_AI_Review3_Faculty_Accurate.pptx"

W = Inches(13.333)
H = Inches(7.5)
NAVY = RGBColor(15, 23, 42)
PANEL = RGBColor(30, 41, 59)
PANEL_ALT = RGBColor(40, 54, 75)
WHITE = RGBColor(248, 250, 252)
MUTED = RGBColor(203, 213, 225)
CYAN = RGBColor(34, 211, 238)
VIOLET = RGBColor(167, 139, 250)
GREEN = RGBColor(52, 211, 153)
AMBER = RGBColor(251, 191, 36)
RED = RGBColor(248, 113, 113)


def add_text(slide, left, top, width, height, text, *, size=18, color=WHITE,
             bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, left, top, width, height, items, *, color=MUTED, size=17):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.space_after = Pt(8)
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.font.bold = False
        paragraph.text = "• " + item
    return box


def card(slide, left, top, width, height, *, accent=CYAN):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = accent
    shape.line.width = Pt(1)
    return shape


def base_slide(prs, title, subtitle="", *, accent=CYAN, number=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.08))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    add_text(slide, Inches(0.55), Inches(0.27), Inches(11.8), Inches(0.45),
             title, size=26, color=WHITE, bold=True)
    if subtitle:
        add_text(slide, Inches(0.57), Inches(0.79), Inches(11.8), Inches(0.3),
                 subtitle, size=11, color=MUTED)
    add_text(slide, Inches(0.55), Inches(7.08), Inches(8), Inches(0.2),
             "SysSense AI  |  Operating Systems Lab  |  Review 3",
             size=9, color=MUTED)
    if number is not None:
        add_text(slide, Inches(12.1), Inches(7.05), Inches(0.65), Inches(0.22),
                 str(number), size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    return slide


def add_table(slide, left, top, width, height, headers, rows, *, accent=CYAN,
              font_size=12):
    table = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height).table
    for col in table.columns:
        col.width = width // len(headers)
    for col_index, header in enumerate(headers):
        cell = table.cell(0, col_index)
        cell.text = str(header)
        cell.fill.solid()
        cell.fill.fore_color.rgb = accent
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL if row_index % 2 else PANEL_ALT
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = WHITE
                    if row == table.rows[0]:
                        run.font.bold = True
    return table


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    band.fill.solid()
    band.fill.fore_color.rgb = NAVY
    band.line.fill.background()
    add_text(slide, Inches(0.8), Inches(1.25), Inches(11.7), Inches(0.9),
             "SysSense AI", size=44, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.0), Inches(2.25), Inches(11.3), Inches(0.8),
             "Live OS Monitoring, Dataset-Trained Forecasting,\nand Evidence-Based Validation",
             size=24, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.0), Inches(3.4), Inches(11.3), Inches(0.4),
             "Project Review 3 | Operating Systems Lab", size=17,
             color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.0), Inches(4.2), Inches(11.3), Inches(0.4),
             "Rohan Kumar (24BDS0463) | Shreyans Bohra (24BCT0013)",
             size=15, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.0), Inches(4.65), Inches(11.3), Inches(0.35),
             "Faculty guide: Dhivyaa C R", size=14, color=MUTED,
             align=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.35),
             "Key principle: live measurements and model estimates are shown separately.",
             size=13, color=GREEN, bold=True, align=PP_ALIGN.CENTER)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    title_slide(prs)

    slide = base_slide(
        prs, "What is live and what is model output?",
        "This distinction makes the demonstration technically accurate.", number=2,
    )
    add_table(slide, Inches(0.55), Inches(1.25), Inches(12.2), Inches(4.9),
              ["Dashboard item", "Source", "Faculty-safe statement"], [
                  ["CPU / RAM / Disk", "Direct psutil snapshot", "Live OS telemetry"],
                  ["History chart", "SQLite collector", "Recorded local telemetry"],
                  ["RF / LSTM CPU-RAM", "Required dataset-trained models", "Experimental nominal 30-second estimates"],
                  ["Disk on Prediction page", "Direct psutil snapshot", "Live value - not forecast"],
                  ["Process signals", "Fixed thresholds", "Rule-based warning - not ML anomaly"],
                  ["Anomaly status", "Isolation Forest", "Experimental signal; evaluate beyond accuracy"],
              ], accent=CYAN, font_size=12)
    add_text(slide, Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.35),
             "The UI labels these categories explicitly to prevent a model estimate from being presented as a measurement.",
             size=13, color=GREEN, align=PP_ALIGN.CENTER)

    slide = base_slide(
        prs, "System architecture and live-data path",
        "Collector cadence is approximately five seconds; forecasts are evaluated against later readings.",
        number=3,
    )
    labels = [
        ("Operating System", "psutil"),
        ("Collector", "CPU / RAM / disk / net / process / thread"),
        ("SQLite", "metrics + process snapshots + forecast log"),
        ("FastAPI", "telemetry, forecast, alerts, evaluation"),
        ("React", "dashboard and presentation UI"),
    ]
    x_positions = [0.45, 2.95, 5.45, 7.95, 10.45]
    for (heading, detail), left in zip(labels, x_positions):
        card(slide, Inches(left), Inches(2.05), Inches(2.35), Inches(2.05), accent=CYAN)
        add_text(slide, Inches(left + 0.12), Inches(2.35), Inches(2.1), Inches(0.35),
                 heading, size=16, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, Inches(left + 0.12), Inches(2.88), Inches(2.1), Inches(0.75),
                 detail, size=12, color=MUTED, align=PP_ALIGN.CENTER)
    for left in [2.55, 5.05, 7.55, 10.05]:
        add_text(slide, Inches(left), Inches(2.83), Inches(0.35), Inches(0.35),
                 ">", size=24, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_bullets(slide, Inches(0.8), Inches(4.7), Inches(11.7), Inches(1.4), [
        "Short positive CPU sampling prevents psutil's first non-blocking 0.0% sentinel from appearing on the dashboard.",
        "Each forecast preserves its source collector sample and is later matched to a real observation.",
    ], size=15)

    slide = base_slide(
        prs, "Dataset requirement and 24-feature contract",
        "The faculty-required dataset is retained; deployment uses only fields the collector really has.",
        accent=VIOLET, number=4,
    )
    card(slide, Inches(0.55), Inches(1.35), Inches(5.8), Inches(4.85), accent=VIOLET)
    add_text(slide, Inches(0.85), Inches(1.65), Inches(5.2), Inches(0.35),
             "Training data", size=19, color=VIOLET, bold=True)
    add_bullets(slide, Inches(0.85), Inches(2.15), Inches(5.1), Inches(3.5), [
        "Mapped Kaggle IT System Performance & Resource Metrics dataset",
        "10,000 benchmark rows",
        "Six rows ahead = nominal 30-second target under mapped five-second timestamps",
        "Chronological splits and embargoes avoid target overlap",
        "The dataset is a benchmark, not a trace from this demo computer",
    ], size=15)
    card(slide, Inches(6.9), Inches(1.35), Inches(5.8), Inches(4.85), accent=CYAN)
    add_text(slide, Inches(7.2), Inches(1.65), Inches(5.2), Inches(0.35),
             "Live-compatible inputs (24)", size=19, color=CYAN, bold=True)
    add_bullets(slide, Inches(7.2), Inches(2.15), Inches(5.1), Inches(3.7), [
        "Current CPU, RAM, disk",
        "CPU/RAM lags, rolling means, deltas, and variation",
        "Network counter rates",
        "Thread count, process count, hour, minute",
        "No fabricated values: incompatible artifacts are disabled",
    ], size=15)

    slide = base_slide(
        prs, "Random Forest and LSTM benchmark evaluation",
        "Forecasting is regression: MAE, RMSE, and R2 are appropriate - not a generic accuracy percentage.",
        accent=GREEN, number=5,
    )
    add_table(slide, Inches(0.55), Inches(1.2), Inches(12.2), Inches(2.3),
              ["Model", "Target", "MAE", "RMSE", "R2", "Persistence MAE"], [
                  ["Random Forest", "CPU", "23.0233", "26.8685", "-0.0052", "31.3095"],
                  ["Random Forest", "RAM", "22.6682", "26.2040", "-0.0123", "30.3470"],
                  ["LSTM", "CPU", "22.9939", "26.8115", "-0.0008", "31.2902"],
                  ["LSTM", "RAM", "22.5530", "26.0690", "-0.0022", "30.3418"],
              ], accent=GREEN, font_size=13)
    card(slide, Inches(0.75), Inches(4.15), Inches(11.8), Inches(1.55), accent=AMBER)
    add_text(slide, Inches(1.0), Inches(4.43), Inches(11.2), Inches(0.3),
             "Interpretation", size=18, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.05), Inches(4.9), Inches(11.1), Inches(0.48),
             "Both models lower MAE versus persistence on the benchmark, but negative R2 means the results do not prove reliable operational forecasting.",
             size=15, color=WHITE, align=PP_ALIGN.CENTER)

    slide = base_slide(
        prs, "Anomaly detection: why headline accuracy is misleading",
        "The anomaly dataset is imbalanced; present the complete held-out evidence.",
        accent=RED, number=6,
    )
    add_table(slide, Inches(0.7), Inches(1.28), Inches(6.0), Inches(4.2),
              ["Held-out metric", "Result"], [
                  ["Accuracy", "98.15%"],
                  ["Balanced accuracy", "58.88%"],
                  ["Precision", "25.00%"],
                  ["Recall", "18.52%"],
                  ["F1", "21.28%"],
                  ["All-normal baseline accuracy", "98.65%"],
              ], accent=RED, font_size=14)
    card(slide, Inches(7.15), Inches(1.28), Inches(5.5), Inches(4.2), accent=AMBER)
    add_text(slide, Inches(7.45), Inches(1.65), Inches(4.9), Inches(0.4),
             "Correct viva statement", size=20, color=AMBER, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(7.55), Inches(2.35), Inches(4.7), Inches(1.95),
             "“We do not claim 98.15% accurate anomaly detection. The all-normal baseline is higher, recall is 18.52%, and F1 is 21.28%. We use the model as an experimental signal and retain transparent threshold rules.”",
             size=17, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.9), Inches(5.95), Inches(11.7), Inches(0.3),
             "Windows System Idle Process is excluded from process rules so idle time is not presented as a false high-CPU anomaly.",
             size=13, color=GREEN, align=PP_ALIGN.CENTER)

    slide = base_slide(
        prs, "Measured local validation after deployment",
        "The dashboard measures live error instead of assuming the benchmark result transfers to this machine.",
        accent=CYAN, number=7,
    )
    steps = [
        ("1", "Forecast", "RF, LSTM, and persistence values are created from a fresh collector snapshot."),
        ("2", "Log", "The source sample, target time, and predicted CPU/RAM are stored in forecast_log."),
        ("3", "Match", "After 30 seconds, the first later collector sample becomes the actual observation."),
        ("4", "Evaluate", "/predict/evaluation reports MAE, RMSE, R2, and +/-5 percentage-point rate."),
    ]
    for index, (number, heading, detail) in enumerate(steps):
        left = 0.6 + index * 3.15
        card(slide, Inches(left), Inches(1.8), Inches(2.75), Inches(3.6), accent=CYAN)
        add_text(slide, Inches(left + 0.85), Inches(2.05), Inches(1.05), Inches(0.55),
                 number, size=30, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, Inches(left + 0.2), Inches(2.85), Inches(2.35), Inches(0.35),
                 heading, size=17, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, Inches(left + 0.2), Inches(3.4), Inches(2.35), Inches(1.25),
                 detail, size=12, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.9), Inches(6.1), Inches(11.5), Inches(0.35),
             "Wait for at least 30 matched forecasts before comparing local models with the persistence baseline.",
             size=14, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

    slide = base_slide(
        prs, "Live demonstration plan",
        "Show evidence in the order that makes data provenance clear.",
        accent=VIOLET, number=8,
    )
    add_bullets(slide, Inches(0.85), Inches(1.35), Inches(11.6), Inches(4.9), [
        "Start the collector; wait about one minute for sufficient history.",
        "Dashboard: point out that CPU, RAM, disk, and network cards are live psutil telemetry.",
        "Process Manager: explain PID, status, CPU, memory, and the difference between process rules and ML.",
        "Prediction: show RF/LSTM estimates as experimental and disk as live-only.",
        "After the horizon: open Live Forecast Evaluation and show matched forecast error and the persistence baseline.",
        "Alerts: show system/process threshold rules separately from the experimental Isolation Forest signal.",
        "API docs: demonstrate /metrics/current, /predict, and /predict/evaluation.",
    ], size=16)
    card(slide, Inches(1.0), Inches(6.0), Inches(11.3), Inches(0.58), accent=GREEN)
    add_text(slide, Inches(1.15), Inches(6.15), Inches(11.0), Inches(0.22),
             "Recommended wording: “We show direct live measurements and separately validate forecasts after they mature.”",
             size=13, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

    slide = base_slide(
        prs, "Limitations and next steps",
        "An accurate presentation states what has been demonstrated and what remains to be validated.",
        accent=AMBER, number=9,
    )
    add_bullets(slide, Inches(0.85), Inches(1.35), Inches(11.6), Inches(4.7), [
        "The Kaggle-derived benchmark is not a local workstation trace; save its exact version and licence.",
        "Benchmark results do not automatically transfer to the demonstration machine.",
        "Collect labelled local workloads and retain an unseen local test period before local retraining or calibration.",
        "Use live matched forecasts to decide whether RF/LSTM beats persistence on this machine.",
        "Improve anomaly labels and tune for precision/recall rather than headline accuracy.",
        "Keep final report, UI labels, and slides synchronized with the 24-feature contract.",
    ], size=17)
    card(slide, Inches(1.35), Inches(6.05), Inches(10.6), Inches(0.5), accent=AMBER)
    add_text(slide, Inches(1.55), Inches(6.17), Inches(10.2), Inches(0.2),
             "Result: a transparent, reproducible OS-monitoring prototype with measurable validation.",
             size=13, color=AMBER, bold=True, align=PP_ALIGN.CENTER)

    slide = base_slide(
        prs, "Thank you",
        "Questions and live demonstration", accent=CYAN, number=10,
    )
    add_text(slide, Inches(1.0), Inches(2.1), Inches(11.3), Inches(0.85),
             "SysSense AI", size=38, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.0), Inches(3.15), Inches(11.3), Inches(0.6),
             "Live monitoring + dataset-trained models + measured validation",
             size=22, color=CYAN, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.0), Inches(4.45), Inches(11.3), Inches(0.35),
             "Documentation: README.md | docs/REVIEW_3_PREPARATION.md",
             size=14, color=MUTED, align=PP_ALIGN.CENTER)

    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
